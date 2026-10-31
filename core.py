# -*- coding: utf-8 -*-
"""Logique métier : base SQLite, import du journal, stocks/CMUP, immobilisations, exports."""
import os
import re
import csv
import sqlite3
import datetime
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Table, TableStyle,
                                PageBreak, Spacer)

# ----------------------------------------------------------------------------
# Utilitaires
# ----------------------------------------------------------------------------


def to_num(v):
    """Convertit '1 234,56' / '(120.5)' / 12 / None en float."""
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\u00a0", "").replace("\u202f", "").replace(" ", "")
    if not s:
        return 0.0
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    if s.startswith("-"):
        neg, s = True, s[1:]
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        x = float(s)
    except ValueError:
        return 0.0
    return -x if neg else x


def fmt(x):
    if x is None:
        return ""
    return f"{x:,.2f}".replace(",", " ").replace(".", ",")


def fmtq(x):
    if x is None or x == 0:
        return ""
    s = f"{x:,.3f}".rstrip("0").rstrip(".")
    return s.replace(",", " ").replace(".", ",") if "." in s or "," in s else s


def parse_date(v):
    """Retourne une date ISO (AAAA-MM-JJ) ou ''."""
    if v is None or v == "":
        return ""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        return (datetime.date(1899, 12, 30) + datetime.timedelta(days=int(v))).isoformat()
    s = str(v).strip()
    for f in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%y", "%d.%m.%Y",
              "%Y/%m/%d", "%d%m%Y", "%d%m%y"):
        try:
            return datetime.datetime.strptime(s, f).date().isoformat()
        except ValueError:
            pass
    try:
        return datetime.datetime.fromisoformat(s).date().isoformat()
    except ValueError:
        return ""


def dfmt(iso):
    if iso and re.match(r"^\d{4}-\d{2}-\d{2}$", str(iso)):
        y, m, d = str(iso).split("-")
        return f"{d}/{m}/{y}"
    return iso or ""


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode()
    return s.lower().strip()


def open_file(path):
    import sys
    import subprocess
    try:
        if sys.platform.startswith("win"):
            os.startfile(path)  # noqa
        elif sys.platform == "darwin":
            subprocess.call(["open", path])
        else:
            subprocess.call(["xdg-open", path])
    except Exception:
        pass


# ----------------------------------------------------------------------------
# Lecture / détection du journal
# ----------------------------------------------------------------------------

def read_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    if ext in (".csv", ".txt"):
        raw = open(path, "rb").read()
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        first = text.splitlines()[0] if text else ""
        delim = max([";", ",", "\t", "|"], key=first.count)
        return [r for r in csv.reader(text.splitlines(), delimiter=delim)]
    raise ValueError("Format non pris en charge. Utilisez .xlsx ou .csv "
                     "(pour un .xls, faites Enregistrer sous > .xlsx).")


def detect_columns(rows):
    for hi, row in enumerate(rows[:40]):
        heads = [norm(c) for c in row]
        used = set()

        def find(kws, excl=()):
            for j, h in enumerate(heads):
                if j in used or not h:
                    continue
                if any(k in h for k in kws) and not any(x in h for x in excl):
                    used.add(j)
                    return j
            return None

        m = {}
        m["debit"] = find(["debit"])
        m["credit"] = find(["credit"])
        m["compte"] = find(["compte", "account", "cpt"], ("libelle", "intitule", "lib", "tiers"))
        m["date"] = find(["date"], ("echeance",))
        m["journal"] = find(["journal", "jnl", "code j"])
        m["piece"] = find(["piece", "n° p", "num", "ref", "document"])
        m["libelle"] = find(["libelle", "designation", "description", "intitule", "objet"], ("compte",))
        if m["libelle"] is None:
            m["libelle"] = find(["libelle", "designation", "description", "intitule", "objet"])
        if m["compte"] is not None and m["debit"] is not None and m["credit"] is not None:
            return hi, m
    raise ValueError("Colonnes introuvables. Le fichier doit contenir au minimum les "
                     "colonnes : Compte, Débit, Crédit (et idéalement Date, Journal, "
                     "Pièce, Libellé) sur une ligne d'en-tête.")


def parse_journal(rows):
    hi, m = detect_columns(rows)

    def g(row, key):
        j = m.get(key)
        return row[j] if j is not None and j < len(row) else ""

    out = []
    for row in rows[hi + 1:]:
        cpt = g(row, "compte")
        if isinstance(cpt, float) and cpt.is_integer():
            cpt = int(cpt)
        cpt = re.sub(r"\s+", "", str(cpt or ""))
        cpt = re.sub(r"\.0$", "", cpt)
        if not re.match(r"^\d{2,}", cpt):
            continue
        d = g(row, "date")
        out.append(dict(
            date=parse_date(d) or (str(d) if d else ""),
            journal=str(g(row, "journal") or "").strip(),
            piece=str(g(row, "piece") or "").strip(),
            compte=cpt,
            libelle=str(g(row, "libelle") or "").strip(),
            debit=to_num(g(row, "debit")),
            credit=to_num(g(row, "credit")),
        ))
    return out


# ----------------------------------------------------------------------------
# Base de données
# ----------------------------------------------------------------------------

SCHEMA = """
CREATE TABLE IF NOT EXISTS settings(k TEXT PRIMARY KEY, v TEXT);
CREATE TABLE IF NOT EXISTS journal(id INTEGER PRIMARY KEY, date TEXT, journal TEXT,
    piece TEXT, compte TEXT, libelle TEXT, debit REAL, credit REAL);
CREATE TABLE IF NOT EXISTS stock_lines(id INTEGER PRIMARY KEY, compte_achat TEXT,
    compte_stock TEXT, date TEXT, piece TEXT, libelle TEXT, article TEXT,
    montant REAL, quantite REAL DEFAULT 0);
CREATE TABLE IF NOT EXISTS stock_articles(compte_stock TEXT, article TEXT,
    unite TEXT DEFAULT '', qte_init REAL DEFAULT 0, val_init REAL DEFAULT 0,
    qte_inv REAL, PRIMARY KEY(compte_stock, article));
CREATE TABLE IF NOT EXISTS stock_sorties(id INTEGER PRIMARY KEY, compte_vente TEXT,
    compte_stock TEXT, date TEXT, piece TEXT, libelle TEXT, article TEXT,
    montant REAL, quantite REAL DEFAULT 0);
CREATE TABLE IF NOT EXISTS immos(id INTEGER PRIMARY KEY, compte TEXT, date TEXT,
    piece TEXT, libelle TEXT, montant REAL, reference TEXT);
CREATE TABLE IF NOT EXISTS composantes(id INTEGER PRIMARY KEY, immo_id INTEGER,
    ref TEXT, sous_ref TEXT, designation TEXT, quantite REAL DEFAULT 1,
    valeur_unitaire REAL DEFAULT 0, duree REAL DEFAULT 0, date_service TEXT,
    localisation TEXT);
"""


def stock_account(compte):
    """601xxx -> 311xxx (marchandises) ; 602xxx -> 321xxx (matières) : même longueur de compte."""
    for p, n in (("601", "311"), ("602", "321")):
        if compte.startswith(p):
            return n + compte[3:]
    return None


def stock_account_vente(compte):
    """701xxx -> 311xxx ; 702xxx -> 321xxx (modifiable ensuite dans le logiciel)."""
    for p, n in (("701", "311"), ("702", "321")):
        if compte.startswith(p):
            return n + compte[3:]
    return None


_STOP = {"achat", "achats", "vente", "ventes", "facture", "fa", "fv", "de", "du", "la", "le",
         "les", "des", "et", "d", "l", "sur", "pour", "un", "une"}


def _tokens(txt):
    return {t for t in re.findall(r"[a-z0-9]+", norm(txt)) if t not in _STOP}


class DB:
    def __init__(self, path):
        self.path = path
        self.c = sqlite3.connect(path)
        self.c.row_factory = sqlite3.Row
        self.c.executescript(SCHEMA)
        self.c.commit()

    # -- paramètres
    def setting(self, k, default=""):
        r = self.c.execute("SELECT v FROM settings WHERE k=?", (k,)).fetchone()
        return r["v"] if r else default

    def set_setting(self, k, v):
        self.c.execute("INSERT OR REPLACE INTO settings(k,v) VALUES(?,?)", (k, v))
        self.c.commit()

    def count(self, table):
        return self.c.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]

    # -- journal
    def import_journal(self, entries):
        self.c.execute("DELETE FROM journal")
        self.c.executemany(
            "INSERT INTO journal(date,journal,piece,compte,libelle,debit,credit) "
            "VALUES(:date,:journal,:piece,:compte,:libelle,:debit,:credit)", entries)
        self.c.commit()

    # -- génération stocks + sorties + immos (conserve les saisies existantes)
    def generate(self):
        c = self.c

        def key(r, cpt):
            return (r[cpt], r["date"], r["piece"], r["libelle"], round(r["montant"] or 0, 2))

        old_lines, old_immos, old_sorties = {}, {}, {}
        for r in c.execute("SELECT * FROM stock_lines ORDER BY id"):
            old_lines.setdefault(key(r, "compte_achat"), []).append(r)
        for r in c.execute("SELECT * FROM immos ORDER BY id"):
            old_immos.setdefault(key(r, "compte"), []).append(r)
        for r in c.execute("SELECT * FROM stock_sorties ORDER BY id"):
            old_sorties.setdefault(key(r, "compte_vente"), []).append(r)
        used_refs = set()

        c.execute("DELETE FROM stock_lines")
        c.execute("DELETE FROM stock_sorties")
        c.execute("DELETE FROM immos")
        n_stock = n_immo = n_new = n_sort = 0
        seq = {}
        sales = []

        for j in c.execute("SELECT * FROM journal ORDER BY id").fetchall():
            cpt = j["compte"]
            montant = round((j["debit"] or 0) - (j["credit"] or 0), 2)
            if cpt.startswith(("601", "602")):
                k = (cpt, j["date"], j["piece"], j["libelle"], montant)
                old = old_lines.get(k, [])
                o = old.pop(0) if old else None
                c.execute(
                    "INSERT INTO stock_lines(compte_achat,compte_stock,date,piece,libelle,"
                    "article,montant,quantite) VALUES(?,?,?,?,?,?,?,?)",
                    (cpt, o["compte_stock"] if o else stock_account(cpt), j["date"], j["piece"],
                     j["libelle"], o["article"] if o else (j["libelle"] or cpt), montant,
                     o["quantite"] if o else 0))
                n_stock += 1
            elif cpt.startswith(("701", "702")):
                sales.append(j)
            elif cpt.startswith(("21", "22", "23", "24")):
                k = (cpt, j["date"], j["piece"], j["libelle"], montant)
                old = old_immos.get(k, [])
                o = old.pop(0) if old else None
                if o:
                    ref = o["reference"]
                else:
                    seq[cpt] = seq.get(cpt, 0)
                    while True:
                        seq[cpt] += 1
                        ref = f"{cpt}-{seq[cpt]:04d}"
                        if ref not in used_refs:
                            break
                used_refs.add(ref)
                cur = c.execute(
                    "INSERT INTO immos(compte,date,piece,libelle,montant,reference) "
                    "VALUES(?,?,?,?,?,?)", (cpt, j["date"], j["piece"], j["libelle"], montant, ref))
                new_id = cur.lastrowid
                n_immo += 1
                if o:
                    c.execute("UPDATE composantes SET immo_id=? WHERE immo_id=?", (new_id, o["id"]))
                elif montant > 0:
                    c.execute(
                        "INSERT INTO composantes(immo_id,ref,sous_ref,designation,quantite,"
                        "valeur_unitaire,duree,date_service,localisation) VALUES(?,?,?,?,?,?,?,?,?)",
                        (new_id, ref + "-C01", "", j["libelle"], 1, montant, 0, j["date"], ""))
                    n_new += 1

        # ventes 701/702 -> sorties de stock (montant = crédit - débit)
        for j in sales:
            cpt = j["compte"]
            m = round((j["credit"] or 0) - (j["debit"] or 0), 2)
            k = (cpt, j["date"], j["piece"], j["libelle"], m)
            old = old_sorties.get(k, [])
            o = old.pop(0) if old else None
            c.execute(
                "INSERT INTO stock_sorties(compte_vente,compte_stock,date,piece,libelle,article,"
                "montant,quantite) VALUES(?,?,?,?,?,?,?,?)",
                (cpt, o["compte_stock"] if o else stock_account_vente(cpt), j["date"], j["piece"],
                 j["libelle"], o["article"] if o else (j["libelle"] or cpt), m,
                 o["quantite"] if o else 0))
            n_sort += 1

        c.execute("DELETE FROM composantes WHERE immo_id NOT IN (SELECT id FROM immos)")
        c.commit()
        self.auto_match()
        return n_stock, n_immo, n_new, n_sort

    # -- lignes de stock
    def set_line_qty(self, lid, q):
        self.c.execute("UPDATE stock_lines SET quantite=? WHERE id=?", (q, lid))
        self.c.commit()

    def set_line_stock_account(self, ids, cpt):
        for i in ids:
            self.c.execute("UPDATE stock_lines SET compte_stock=? WHERE id=?", (cpt, i))
        self.c.commit()

    def set_line_article(self, ids, new):
        new = new.strip()
        if not new:
            return
        for i in ids:
            r = self.c.execute("SELECT compte_stock, article FROM stock_lines WHERE id=?", (i,)).fetchone()
            if not r:
                continue
            self.c.execute("UPDATE stock_lines SET article=? WHERE id=?", (new, i))
            left = self.c.execute("SELECT COUNT(*) FROM stock_lines WHERE compte_stock=? AND article=?",
                                  (r["compte_stock"], r["article"])).fetchone()[0]
            if left == 0:
                self.c.execute("UPDATE stock_sorties SET article=? WHERE compte_stock=? AND article=?",
                               (new, r["compte_stock"], r["article"]))
                try:
                    self.c.execute("UPDATE stock_articles SET article=? WHERE compte_stock=? AND article=?",
                                   (new, r["compte_stock"], r["article"]))
                except sqlite3.IntegrityError:
                    self.c.execute("DELETE FROM stock_articles WHERE compte_stock=? AND article=?",
                                   (r["compte_stock"], r["article"]))
        self.c.commit()

    def set_article_field(self, compte, article, field, value):
        assert field in ("unite", "qte_init", "val_init", "qte_inv")
        self.c.execute("INSERT OR IGNORE INTO stock_articles(compte_stock, article) VALUES(?,?)",
                       (compte, article))
        self.c.execute(f"UPDATE stock_articles SET {field}=? WHERE compte_stock=? AND article=?",
                       (value, compte, article))
        self.c.commit()

    def fiches(self):
        c = self.c
        arts = {(r["compte_stock"], r["article"]): r for r in c.execute("SELECT * FROM stock_articles")}
        ach, sor = {}, {}
        for r in c.execute("SELECT compte_stock, article, SUM(montant) m, SUM(quantite) q, COUNT(*) n, "
                           "SUM(CASE WHEN COALESCE(quantite,0)=0 THEN 1 ELSE 0 END) nq "
                           "FROM stock_lines GROUP BY compte_stock, article"):
            ach[(r["compte_stock"], r["article"])] = (r["m"] or 0.0, r["q"] or 0.0, r["n"], r["nq"] or 0)
        for r in c.execute("SELECT compte_stock, article, SUM(montant) m, SUM(quantite) q, COUNT(*) n "
                           "FROM stock_sorties GROUP BY compte_stock, article"):
            sor[(r["compte_stock"], r["article"])] = (r["m"] or 0.0, r["q"] or 0.0, r["n"])
        out = []
        for key in sorted(set(ach) | set(sor)):
            a = arts.get(key)
            m_ach, q_ach, nb, nq = ach.get(key, (0.0, 0.0, 0, 0))
            vente, q_sort, nbs = sor.get(key, (0.0, 0.0, 0))
            qi = (a["qte_init"] or 0.0) if a else 0.0
            vi = (a["val_init"] or 0.0) if a else 0.0
            qinv = a["qte_inv"] if a else None
            q_tot = qi + q_ach
            v_tot = vi + m_ach
            cmup = v_tot / q_tot if q_tot else 0.0
            q_theo = q_tot - q_sort
            counted = qinv is not None
            q_final = qinv if counted else q_theo
            ecart_q = (qinv - q_theo) if counted else 0.0
            out.append(dict(
                compte=key[0], article=key[1], unite=(a["unite"] or "") if a else "",
                qte_init=qi, val_init=vi, q_ach=q_ach, m_ach=m_ach, q_tot=q_tot, v_tot=v_tot,
                cmup=cmup, q_sort=q_sort, vente=vente, v_sort=q_sort * cmup, q_theo=q_theo,
                qte_inv=qinv, counted=counted, ecart_q=ecart_q, ecart_v=ecart_q * cmup,
                q_final=q_final, v_inv=q_final * cmup, nb=nb, nbs=nbs, sans_qte=nq, cmup_src="article"))
        # CMUP de repli : article sans achat -> CMUP moyen du compte de stock, puis de la classe (311, 321...)
        agg_c, agg_k = {}, {}
        for f in out:
            if f["q_tot"] > 0:
                for d, k in ((agg_c, f["compte"]), (agg_k, f["compte"][:3])):
                    x = d.setdefault(k, [0.0, 0.0])
                    x[0] += f["v_tot"]
                    x[1] += f["q_tot"]
        for f in out:
            if f["q_tot"] > 0:
                continue
            if f["compte"] in agg_c:
                v, q = agg_c[f["compte"]]
                src = "compte"
            elif f["compte"][:3] in agg_k:
                v, q = agg_k[f["compte"][:3]]
                src = "classe"
            else:
                v, q, src = 0.0, 1.0, "aucun"
            f["cmup"] = v / q
            f["cmup_src"] = src
            f["v_sort"] = f["q_sort"] * f["cmup"]
            f["ecart_v"] = f["ecart_q"] * f["cmup"]
            f["v_inv"] = f["q_final"] * f["cmup"]
        return out

    def ecritures(self):
        """Écritures d'inventaire proposées (variation de stocks 603x)."""
        agg = {}
        for f in self.fiches():
            a = agg.setdefault(f["compte"], [0.0, 0.0])
            a[0] += f["val_init"]
            a[1] += f["v_inv"]
        out = []
        for cpt, (vi, vf) in sorted(agg.items()):
            var = "603" + (cpt[1:2] if len(cpt) > 1 else "")
            vi, vf = round(vi, 2), round(vf, 2)
            if vi:
                out += [(var, f"Annulation stock initial {cpt}", vi, 0), (cpt, "Annulation stock initial", 0, vi)]
            if vf:
                out += [(cpt, "Constatation stock final", vf, 0), (var, f"Constatation stock final {cpt}", 0, vf)]
        return out

    # -- sorties de stock (ventes 701/702)
    def set_sortie_qty(self, sid, q):
        self.c.execute("UPDATE stock_sorties SET quantite=? WHERE id=?", (q, sid))
        self.c.commit()

    def set_sortie_article(self, ids, art):
        art = art.strip()
        if art:
            for i in ids:
                self.c.execute("UPDATE stock_sorties SET article=? WHERE id=?", (art, i))
            self.c.commit()

    def set_sortie_stock_account(self, ids, cpt):
        for i in ids:
            self.c.execute("UPDATE stock_sorties SET compte_stock=? WHERE id=?", (cpt, i))
        self.c.commit()

    def set_sortie_link(self, ids, compte_stock, article):
        for i in ids:
            self.c.execute("UPDATE stock_sorties SET compte_stock=?, article=? WHERE id=?",
                           (compte_stock, article, i))
        self.c.commit()

    def auto_match(self):
        """Rattache les ventes non reliées à l'article d'achat le plus proche (par mots-clés)."""
        c = self.c
        cands = [(r[0], r[1]) for r in c.execute("SELECT DISTINCT compte_stock, article FROM stock_lines")]
        known = set(cands)
        n = 0
        for sd in c.execute("SELECT * FROM stock_sorties").fetchall():
            if (sd["compte_stock"], sd["article"]) in known:
                continue
            ts = _tokens(sd["article"] or "") | _tokens(sd["libelle"] or "")
            best, best_sc = None, 0.0
            for cs, art in cands:
                ta = _tokens(art)
                if not ta or not ts:
                    continue
                ov = len(ta & ts)
                sc = ov / min(len(ta), len(ts))
                if ov >= min(2, len(ta)) and sc >= 0.6:
                    sc += 0.01 if cs == sd["compte_stock"] else 0
                    if sc > best_sc:
                        best, best_sc = (cs, art), sc
            if not best:
                same = [ar for cs, ar in cands if cs == sd["compte_stock"]]
                if len(same) == 1:
                    best = (sd["compte_stock"], same[0])
            if best:
                c.execute("UPDATE stock_sorties SET compte_stock=?, article=? WHERE id=?",
                          (best[0], best[1], sd["id"]))
                n += 1
        c.commit()
        return n

    def sorties_rows(self):
        cm = {(f["compte"], f["article"]): (f["cmup"], f["cmup_src"]) for f in self.fiches()}
        out = []
        for r in self.c.execute("SELECT * FROM stock_sorties ORDER BY compte_stock, article, date, id"):
            cmup, src = cm.get((r["compte_stock"], r["article"]), (0.0, "aucun"))
            cout = (r["quantite"] or 0) * cmup
            d = dict(r)
            d.update(cmup=cmup, cmup_src=src, cout=cout, marge=(r["montant"] or 0) - cout)
            out.append(d)
        return out

    # -- immobilisations
    def comps(self, immo_id):
        return self.c.execute("SELECT * FROM composantes WHERE immo_id=? ORDER BY ref, sous_ref, id",
                              (immo_id,)).fetchall()

    def add_comp(self, immo_id):
        im = self.c.execute("SELECT * FROM immos WHERE id=?", (immo_id,)).fetchone()
        n = len(self.comps(immo_id)) + 1
        ref = f"{im['reference']}-C{n:02d}"
        self.c.execute(
            "INSERT INTO composantes(immo_id,ref,sous_ref,designation,quantite,valeur_unitaire,"
            "duree,date_service,localisation) VALUES(?,?,?,?,?,?,?,?,?)",
            (immo_id, ref, "", "", 1, 0, 0, im["date"], ""))
        self.c.commit()

    def set_comp(self, cid, field, value):
        assert field in ("ref", "sous_ref", "designation", "quantite", "valeur_unitaire",
                         "duree", "date_service", "localisation")
        self.c.execute(f"UPDATE composantes SET {field}=? WHERE id=?", (value, cid))
        self.c.commit()

    def del_comps(self, ids):
        for i in ids:
            self.c.execute("DELETE FROM composantes WHERE id=?", (i,))
        self.c.commit()

    def dup_comp(self, cid):
        r = self.c.execute("SELECT * FROM composantes WHERE id=?", (cid,)).fetchone()
        self.c.execute(
            "INSERT INTO composantes(immo_id,ref,sous_ref,designation,quantite,valeur_unitaire,"
            "duree,date_service,localisation) VALUES(?,?,?,?,?,?,?,?,?)",
            (r["immo_id"], r["ref"], r["sous_ref"], r["designation"], r["quantite"],
             r["valeur_unitaire"], r["duree"], r["date_service"], r["localisation"]))
        self.c.commit()

    def split_comp(self, cid):
        """Éclate une composante de quantité N en N lignes de quantité 1 avec sous-références."""
        r = self.c.execute("SELECT * FROM composantes WHERE id=?", (cid,)).fetchone()
        n = int(round(r["quantite"] or 0))
        if n < 2:
            return 0
        for i in range(1, n + 1):
            self.c.execute(
                "INSERT INTO composantes(immo_id,ref,sous_ref,designation,quantite,valeur_unitaire,"
                "duree,date_service,localisation) VALUES(?,?,?,?,?,?,?,?,?)",
                (r["immo_id"], r["ref"], f"{r['ref']}-{i:03d}", r["designation"], 1,
                 r["valeur_unitaire"], r["duree"], r["date_service"], r["localisation"]))
        self.c.execute("DELETE FROM composantes WHERE id=?", (cid,))
        self.c.commit()
        return n

    def adjust_comp(self, cid):
        """Ajuste la valeur unitaire pour que le total des composantes = montant comptable."""
        r = self.c.execute("SELECT * FROM composantes WHERE id=?", (cid,)).fetchone()
        im = self.c.execute("SELECT * FROM immos WHERE id=?", (r["immo_id"],)).fetchone()
        others = sum((x["quantite"] or 0) * (x["valeur_unitaire"] or 0)
                     for x in self.comps(r["immo_id"]) if x["id"] != cid)
        q = r["quantite"] or 0
        if q:
            self.c.execute("UPDATE composantes SET valeur_unitaire=? WHERE id=?",
                           (round((im["montant"] - others) / q, 4), cid))
            self.c.commit()

    def registre(self, mode="compte", value="", cloture=None):
        cloture = cloture or datetime.date.today()
        q = ("SELECT c.*, i.compte, i.date AS date_acq, i.reference FROM composantes c "
             "JOIN immos i ON i.id=c.immo_id ORDER BY i.compte, c.ref, c.sous_ref, c.id")
        rows = []
        for r in self.c.execute(q):
            if value:
                if mode == "compte" and not r["compte"].startswith(value):
                    continue
                if mode == "composante" and r["ref"] != value:
                    continue
            val = (r["quantite"] or 0) * (r["valeur_unitaire"] or 0)
            duree = r["duree"] or 0
            dot = val / duree if duree else 0.0
            cumul = 0.0
            ds = r["date_service"]
            if duree and re.match(r"^\d{4}-\d{2}-\d{2}$", ds or ""):
                d0 = datetime.date.fromisoformat(ds)
                ratio = min(max((cloture - d0).days / (duree * 365.25), 0), 1)
                cumul = val * ratio
            rows.append(dict(
                compte=r["compte"], reference=r["reference"], ref=r["ref"], sous_ref=r["sous_ref"] or "",
                designation=r["designation"] or "", date_acq=r["date_acq"], date_service=ds or "",
                quantite=r["quantite"] or 0, pu=r["valeur_unitaire"] or 0, valeur=val, duree=duree,
                dotation=dot, cumul=cumul, vnc=val - cumul, localisation=r["localisation"] or ""))
        return rows

    def registre_values(self):
        comptes = [r[0] for r in self.c.execute("SELECT DISTINCT compte FROM immos ORDER BY compte")]
        refs = [r[0] for r in self.c.execute("SELECT DISTINCT ref FROM composantes ORDER BY ref")]
        return comptes, refs


# ----------------------------------------------------------------------------
# Exports Excel
# ----------------------------------------------------------------------------

HFILL = PatternFill("solid", fgColor="1F3864")
HFONT = Font(bold=True, color="FFFFFF")
NUM = '#,##0.00'


def _header(ws, row, titles):
    for j, t in enumerate(titles, 1):
        c = ws.cell(row=row, column=j, value=t)
        c.fill, c.font = HFILL, HFONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, widths):
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


STOCK_TITLES = ["Compte stock", "Article", "Unité", "Qté initiale", "Valeur initiale", "Qté achats",
                "Montant achats", "CMUP", "Qté sorties", "Coût des sorties", "Qté théorique",
                "Qté inventaire", "Écart inventaire (qté)", "Valeur stock final"]


def _stock_vals(f):
    return [f["compte"], f["article"], f["unite"], f["qte_init"], f["val_init"], f["q_ach"],
            f["m_ach"], round(f["cmup"], 4), f["q_sort"], round(f["v_sort"], 2), f["q_theo"],
            f["qte_inv"] if f["counted"] else "", round(f["ecart_q"], 3), round(f["v_inv"], 2)]


def export_stock_xlsx(db, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Récapitulatif"
    ws["A1"] = f"{db.setting('entreprise')} - Inventaire des stocks - Exercice {db.setting('exercice')}"
    ws["A1"].font = Font(bold=True, size=13)
    _header(ws, 3, STOCK_TITLES)
    fiches = db.fiches()
    r = 4
    for f in fiches:
        for j, v in enumerate(_stock_vals(f), 1):
            ws.cell(row=r, column=j, value=v)
        r += 1
    ws.cell(row=r, column=2, value="TOTAL").font = Font(bold=True)
    for col in (5, 7, 10, 14):
        L = get_column_letter(col)
        c = ws.cell(row=r, column=col, value=f"=SUM({L}4:{L}{r - 1})")
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=4, max_row=r, min_col=4, max_col=14):
        for c in row:
            c.number_format = '#,##0.0000' if c.column == 8 else NUM
    _autosize(ws, [13, 38, 8, 12, 16, 12, 16, 12, 12, 16, 14, 14, 16, 18])
    ws.freeze_panes = "A4"

    groups = {}
    for f in fiches:
        groups.setdefault(f["compte"], []).append(f)
    for cpt, items in groups.items():
        w = wb.create_sheet(f"Fiche {cpt}"[:31])
        w["A1"] = f"FICHE D'INVENTAIRE - Compte {cpt}"
        w["A1"].font = Font(bold=True, size=13)
        _header(w, 3, STOCK_TITLES[1:])
        for i, f in enumerate(items, 4):
            for j, v in enumerate(_stock_vals(f)[1:], 1):
                w.cell(row=i, column=j, value=v).number_format = NUM
        _autosize(w, [38, 8, 12, 16, 12, 16, 12, 12, 16, 14, 14, 16, 18])

    wso = wb.create_sheet("Sorties de stock")
    _header(wso, 1, ["Compte stock", "Article", "Date", "Pièce", "Libellé", "Compte vente", "Qté sortie",
                     "Prix de vente", "CMUP", "Coût de sortie", "Marge brute"])
    for i, d in enumerate(db.sorties_rows(), 2):
        vals = [d["compte_stock"], d["article"], dfmt(d["date"]), d["piece"], d["libelle"], d["compte_vente"],
                d["quantite"], d["montant"], round(d["cmup"], 4), round(d["cout"], 2), round(d["marge"], 2)]
        for j, v in enumerate(vals, 1):
            c = wso.cell(row=i, column=j, value=v)
            if j >= 7:
                c.number_format = NUM
    _autosize(wso, [13, 34, 11, 12, 38, 13, 12, 16, 12, 16, 16])

    we = wb.create_sheet("Écritures proposées")
    _header(we, 1, ["Compte", "Libellé", "Débit", "Crédit"])
    for i, e in enumerate(db.ecritures(), 2):
        for j, v in enumerate(e, 1):
            we.cell(row=i, column=j, value=v)
            if j > 2:
                we.cell(row=i, column=j).number_format = NUM
    _autosize(we, [12, 45, 16, 16])
    wb.save(path)


REG_COLS = [("Compte", "compte", 1.4), ("Réf. immo", "reference", 2.0), ("Réf. composante", "ref", 2.2),
            ("Sous-réf.", "sous_ref", 1.8), ("Désignation", "designation", 3.6),
            ("Date acq.", "date_acq", 1.7), ("Mise en service", "date_service", 1.7),
            ("Qté", "quantite", 1.0), ("P.U.", "pu", 1.8), ("Valeur brute", "valeur", 2.0),
            ("Durée (ans)", "duree", 1.0), ("Dotation annuelle", "dotation", 1.9),
            ("Cumul amort.", "cumul", 1.9), ("VNC", "vnc", 1.9), ("Localisation", "localisation", 1.8)]
REG_NUM = {"quantite", "pu", "valeur", "duree", "dotation", "cumul", "vnc"}
REG_SUM = ("valeur", "dotation", "cumul", "vnc")


def export_registre_xlsx(db, path, mode, value, cloture):
    rows = db.registre(mode, value, cloture)
    wb = Workbook()
    ws = wb.active
    ws.title = "Registre"
    ws["A1"] = (f"{db.setting('entreprise')} - REGISTRE DES IMMOBILISATIONS - Situation au "
                f"{cloture.strftime('%d/%m/%Y')}")
    ws["A1"].font = Font(bold=True, size=13)
    _header(ws, 3, [c[0] for c in REG_COLS])
    r = 4
    groups = []
    if mode == "compte":
        for row in rows:
            if not groups or groups[-1][0] != row["compte"]:
                groups.append((row["compte"], []))
            groups[-1][1].append(row)
    else:
        groups = [("", rows)]
    tot = {k: 0.0 for k in REG_SUM}
    for g, items in groups:
        if g:
            ws.cell(row=r, column=1, value=f"Compte {g}").font = Font(bold=True)
            r += 1
        sub = {k: 0.0 for k in REG_SUM}
        for row in items:
            for j, (_, key, _) in enumerate(REG_COLS, 1):
                v = row[key]
                if key in ("date_acq", "date_service"):
                    v = dfmt(v)
                c = ws.cell(row=r, column=j, value=v)
                if key in REG_NUM:
                    c.number_format = NUM
            for k in REG_SUM:
                sub[k] += row[k]
                tot[k] += row[k]
            r += 1
        if g:
            ws.cell(row=r, column=5, value=f"Sous-total {g}").font = Font(bold=True)
            for k in REG_SUM:
                j = [c[1] for c in REG_COLS].index(k) + 1
                c = ws.cell(row=r, column=j, value=sub[k])
                c.font, c.number_format = Font(bold=True), NUM
            r += 2
    ws.cell(row=r, column=5, value="TOTAL GÉNÉRAL").font = Font(bold=True)
    for k in REG_SUM:
        j = [c[1] for c in REG_COLS].index(k) + 1
        c = ws.cell(row=r, column=j, value=tot[k])
        c.font, c.number_format = Font(bold=True), NUM
    _autosize(ws, [10, 14, 16, 14, 34, 12, 14, 8, 14, 16, 10, 16, 16, 16, 16])
    ws.freeze_panes = "A4"
    wb.save(path)


# ----------------------------------------------------------------------------
# Exports PDF
# ----------------------------------------------------------------------------

_ST = getSampleStyleSheet()
SM = ParagraphStyle("sm", parent=_ST["Normal"], fontSize=7, leading=8)
SMB = ParagraphStyle("smb", parent=SM, fontName="Helvetica-Bold")
SMH = ParagraphStyle("smh", parent=SM, fontName="Helvetica-Bold", textColor=colors.white, alignment=1)


def _H(row):
    return [Paragraph(str(t), SMH) for t in row]


def _P(t, bold=False):
    t = str(t if t is not None else "").replace("&", "&amp;").replace("<", "&lt;")
    return Paragraph(t, SMB if bold else SM)


def _style(header_rows=1, right_from=3, font=7):
    return TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), font),
        ("BACKGROUND", (0, 0), (-1, header_rows - 1), colors.HexColor("#1F3864")),
        ("TEXTCOLOR", (0, 0), (-1, header_rows - 1), colors.white),
        ("FONTNAME", (0, 0), (-1, header_rows - 1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
        ("ALIGN", (right_from, header_rows), (-1, -1), "RIGHT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ])


def _doc(path, title, db):
    ent, ex = db.setting("entreprise"), db.setting("exercice")
    W, H = landscape(A4)

    def hf(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(1 * cm, H - 0.7 * cm, f"{ent}  -  Exercice {ex}")
        canvas.drawRightString(W - 1 * cm, 0.6 * cm, f"{title}  -  Page {doc.page}")
        canvas.restoreState()

    d = SimpleDocTemplate(path, pagesize=landscape(A4), leftMargin=1 * cm, rightMargin=1 * cm,
                          topMargin=1.2 * cm, bottomMargin=1.1 * cm, title=title)
    return d, hf


def _sign_block():
    t = Table([["Inventaire réalisé par :", "Date :", "Visa du responsable :"], ["", "", ""]],
              colWidths=[9 * cm, 6 * cm, 9 * cm], rowHeights=[0.6 * cm, 1.6 * cm])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
    return t


def export_stock_pdf(db, path, individual=False):
    title = "Fiches individuelles de stock" if individual else "Fiches d'inventaire des stocks"
    doc, hf = _doc(path, title, db)
    story = []
    fiches = db.fiches()
    if not fiches:
        story.append(Paragraph("Aucune ligne de stock.", _ST["Normal"]))
    if not individual:
        groups = {}
        for f in fiches:
            groups.setdefault(f["compte"], []).append(f)
        widths = [0.8, 4.6, 1.2, 1.7, 2.2, 1.7, 2.3, 1.8, 1.7, 2.2, 1.8, 1.8, 1.5, 2.4]
        for cpt, items in groups.items():
            story.append(Paragraph(f"FICHE D'INVENTAIRE DES STOCKS - Compte {cpt}", _ST["Title"]))
            data = [_H(["N°", "Article", "Unité", "Qté init.", "Val. init.", "Qté achats", "Montant achats",
                        "CMUP", "Qté sorties", "Coût des sorties", "Qté théorique", "Qté inventaire",
                        "Écart qté", "Valeur stock final"])]
            tot = [0.0] * 4
            for i, f in enumerate(items, 1):
                data.append([i, _P(f["article"]), f["unite"], fmtq(f["qte_init"]), fmt(f["val_init"]),
                             fmtq(f["q_ach"]), fmt(f["m_ach"]), fmt(f["cmup"]), fmtq(f["q_sort"]),
                             fmt(f["v_sort"]), fmtq(f["q_theo"]) or "0",
                             fmtq(f["qte_inv"]) if f["counted"] else "-", fmtq(f["ecart_q"]),
                             fmt(f["v_inv"])])
                tot[0] += f["val_init"]; tot[1] += f["m_ach"]; tot[2] += f["v_sort"]; tot[3] += f["v_inv"]
            data.append(["", _P("TOTAL", True), "", "", fmt(tot[0]), "", fmt(tot[1]), "", "", fmt(tot[2]),
                         "", "", "", fmt(tot[3])])
            t = Table(data, colWidths=[w * cm for w in widths], repeatRows=1)
            t.setStyle(_style(1, 3, 6.5))
            t.setStyle(TableStyle([("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                                   ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E7ECF5"))]))
            story += [t, Spacer(1, 0.2 * cm),
                      Paragraph("Quantité théorique = stock initial + achats - sorties (ventes). "
                                "« - » : quantité inventoriée non saisie, la quantité théorique est retenue.", SM),
                      Spacer(1, 0.4 * cm), _sign_block(), PageBreak()]
        ec = db.ecritures()
        if ec:
            story.append(Paragraph("ÉCRITURES D'INVENTAIRE PROPOSÉES (variation des stocks)", _ST["Title"]))
            data = [["Compte", "Libellé", "Débit", "Crédit"]]
            data += [[a, b, fmt(c) if c else "", fmt(d) if d else ""] for a, b, c, d in ec]
            t = Table(data, colWidths=[3 * cm, 12 * cm, 4 * cm, 4 * cm], repeatRows=1)
            t.setStyle(_style(1, 2, 8))
            story.append(t)
        elif story and isinstance(story[-1], PageBreak):
            story.pop()
    else:
        for n, f in enumerate(fiches):
            story.append(Paragraph("FICHE INDIVIDUELLE DE STOCK", _ST["Title"]))
            hdr = Table([[f"Compte de stock : {f['compte']}", f"Article : {f['article'][:70]}",
                          f"Unité : {f['unite']}"]], colWidths=[7 * cm, 14 * cm, 6 * cm])
            hdr.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 9), ("BOX", (0, 0), (-1, -1), 0.5, colors.grey)]))
            story += [hdr, Spacer(1, 0.3 * cm), Paragraph("<b>ENTRÉES (stock initial et achats)</b>", _ST["Normal"])]
            data = [["Date", "Pièce", "Libellé", "Compte achat", "Quantité", "Montant", "P.U."]]
            if f["qte_init"] or f["val_init"]:
                pu = f["val_init"] / f["qte_init"] if f["qte_init"] else 0
                data.append(["", "", "STOCK INITIAL", "", fmtq(f["qte_init"]), fmt(f["val_init"]), fmt(pu)])
            for r in db.c.execute("SELECT * FROM stock_lines WHERE compte_stock=? AND article=? ORDER BY date, id",
                                  (f["compte"], f["article"])):
                pu = r["montant"] / r["quantite"] if r["quantite"] else 0
                data.append([dfmt(r["date"]), r["piece"], _P(r["libelle"]), r["compte_achat"],
                             fmtq(r["quantite"]), fmt(r["montant"]), fmt(pu)])
            t = Table(data, colWidths=[2.2 * cm, 3 * cm, 10 * cm, 2.6 * cm, 2.6 * cm, 3.4 * cm, 3.2 * cm], repeatRows=1)
            t.setStyle(_style(1, 4, 8))
            story += [t, Spacer(1, 0.3 * cm)]
            sorties = db.c.execute("SELECT * FROM stock_sorties WHERE compte_stock=? AND article=? ORDER BY date, id",
                                   (f["compte"], f["article"])).fetchall()
            if sorties:
                story.append(Paragraph("<b>SORTIES (ventes) valorisées au CMUP</b>", _ST["Normal"]))
                data = [["Date", "Pièce", "Libellé", "Compte vente", "Qté sortie", "Prix de vente", "Coût (CMUP)"]]
                for r in sorties:
                    data.append([dfmt(r["date"]), r["piece"], _P(r["libelle"]), r["compte_vente"],
                                 fmtq(r["quantite"]), fmt(r["montant"]), fmt((r["quantite"] or 0) * f["cmup"])])
                t = Table(data, colWidths=[2.2 * cm, 3 * cm, 10 * cm, 2.6 * cm, 2.6 * cm, 3.4 * cm, 3.2 * cm], repeatRows=1)
                t.setStyle(_style(1, 4, 8))
                story += [t, Spacer(1, 0.3 * cm)]
            calc = Table([
                ["Quantité totale (stock initial + achats)", fmtq(f["q_tot"])],
                ["Valeur totale (stock initial + achats)", fmt(f["v_tot"])],
                ["CMUP = valeur totale / quantité totale" if f["cmup_src"] == "article" else
                 "CMUP moyen du compte de stock (aucun achat sur cet article)", fmt(f["cmup"])],
                ["Quantité sortie (ventes)", fmtq(f["q_sort"])],
                ["Coût des sorties = quantité sortie x CMUP", fmt(f["v_sort"])],
                ["Quantité théorique en stock", fmtq(f["q_theo"]) or "0"],
                ["Quantité inventoriée (comptage physique)", fmtq(f["qte_inv"]) if f["counted"] else "non saisie (théorique retenue)"],
                ["Écart d'inventaire (inventoriée - théorique)", fmtq(f["ecart_q"]) or "0"],
                ["VALEUR DU STOCK FINAL = quantité retenue x CMUP", fmt(f["v_inv"])],
            ], colWidths=[14 * cm, 5 * cm])
            calc.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                                      ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                                      ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                                      ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E7ECF5"))]))
            story += [calc, Spacer(1, 0.4 * cm), _sign_block()]
            if n < len(fiches) - 1:
                story.append(PageBreak())
    doc.build(story, onFirstPage=hf, onLaterPages=hf)


def export_registre_pdf(db, path, mode, value, cloture):
    rows = db.registre(mode, value, cloture)
    doc, hf = _doc(path, "Registre des immobilisations", db)
    story = []
    sub_title = f"Situation au {cloture.strftime('%d/%m/%Y')}"
    if value:
        sub_title += f" - {'Compte' if mode == 'compte' else 'Composante'} {value}"
    heads = [c[0] for c in REG_COLS]
    widths = [c[2] * cm for c in REG_COLS]
    first_num = [c[1] for c in REG_COLS].index("quantite")

    def table(items, label_total):
        data = [_H(heads)]
        tot = {k: 0.0 for k in REG_SUM}
        for r in items:
            line = []
            for _, key, _ in REG_COLS:
                v = r[key]
                if key in ("date_acq", "date_service"):
                    v = dfmt(v)
                elif key == "quantite":
                    v = fmtq(v)
                elif key == "duree":
                    v = fmtq(v)
                elif key in REG_NUM:
                    v = fmt(v)
                elif key in ("reference", "ref", "sous_ref", "designation", "localisation"):
                    v = _P(v)
                line.append(v)
            data.append(line)
            for k in REG_SUM:
                tot[k] += r[k]
        tl = [""] * len(REG_COLS)
        tl[4] = _P(label_total, True)
        for k in REG_SUM:
            tl[[c[1] for c in REG_COLS].index(k)] = fmt(tot[k])
        data.append(tl)
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(_style(1, first_num, 6.5))
        t.setStyle(TableStyle([("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                               ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#E7ECF5"))]))
        return t

    if not rows:
        story.append(Paragraph("Aucune composante à afficher.", _ST["Normal"]))
    elif mode == "compte":
        groups = []
        for r in rows:
            if not groups or groups[-1][0] != r["compte"]:
                groups.append((r["compte"], []))
            groups[-1][1].append(r)
        for i, (g, items) in enumerate(groups):
            story.append(Paragraph(f"REGISTRE DES IMMOBILISATIONS - Compte {g}", _ST["Title"]))
            story.append(Paragraph(sub_title, _ST["Normal"]))
            story += [Spacer(1, 0.3 * cm), table(items, f"Total compte {g}")]
            if i < len(groups) - 1:
                story.append(PageBreak())
    else:
        story.append(Paragraph("REGISTRE DES IMMOBILISATIONS - par composante", _ST["Title"]))
        story.append(Paragraph(sub_title, _ST["Normal"]))
        rows.sort(key=lambda r: (r["ref"], r["sous_ref"]))
        story += [Spacer(1, 0.3 * cm), table(rows, "TOTAL")]
    doc.build(story, onFirstPage=hf, onLaterPages=hf)


def export_journal_template(path):
    """Modèle Excel de journal à importer (en-têtes + exemples + notice)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Journal"
    heads = ["Date", "Journal", "Pièce", "Compte", "Libellé", "Débit", "Crédit"]
    _header(ws, 1, heads)
    ex = [
        (datetime.date(2025, 1, 5), "AC", "FA001", "601100", "Achat ciment CPJ 35", 1500000, 0),
        (datetime.date(2025, 1, 5), "AC", "FA001", "445200", "TVA récupérable", 270000, 0),
        (datetime.date(2025, 1, 5), "AC", "FA001", "401000", "Fournisseur ABC", 0, 1770000),
        (datetime.date(2025, 4, 10), "AC", "FA004", "602100", "Achat huile moteur", 350000, 0),
        (datetime.date(2025, 5, 15), "VE", "FV001", "701100", "Vente ciment CPJ 35", 0, 2400000),
        (datetime.date(2025, 2, 15), "AC", "FA010", "241000", "Matériel de transport - Camion benne", 25000000, 0),
        (datetime.date(2025, 2, 20), "AC", "FA011", "244100", "Mobilier de bureau - 10 bureaux", 3000000, 0),
    ]
    for i, row in enumerate(ex, 2):
        for j, v in enumerate(row, 1):
            c = ws.cell(row=i, column=j, value=v)
            if j == 1:
                c.number_format = "DD/MM/YYYY"
            elif j in (6, 7):
                c.number_format = NUM
            elif j == 4:
                c.number_format = "@"
                c.value = str(v)
    _autosize(ws, [12, 10, 12, 12, 42, 16, 16])
    ws.freeze_panes = "A2"
    wn = wb.create_sheet("Notice")
    lines = [
        "MODÈLE DE JOURNAL À IMPORTER",
        "",
        "1. Remplacez les lignes d'exemple (feuille « Journal ») par les écritures de votre journal.",
        "2. Conservez la ligne d'en-tête : Date, Journal, Pièce, Compte, Libellé, Débit, Crédit.",
        "3. Colonnes obligatoires : Compte, Débit, Crédit. Les autres sont recommandées.",
        "4. Compte : numéro de compte général (ex. 601100, 2441). Format texte conseillé.",
        "5. Débit / Crédit : montants numériques (laisser 0 ou vide si sans objet).",
        "6. Comptes retenus : 601 et 602 (achats -> stocks), 701 et 702 (ventes -> sorties de stock), 21 à 24 (immobilisations).",
        "7. Les autres comptes (TVA, fournisseurs, etc.) peuvent rester : ils sont ignorés pour la génération.",
        "8. Enregistrez en .xlsx (ou .csv) puis utilisez « Importer le journal ».",
    ]
    for i, t in enumerate(lines, 1):
        wn.cell(row=i, column=1, value=t)
    wn["A1"].font = Font(bold=True, size=13)
    wn.column_dimensions["A"].width = 110
    wb.save(path)
